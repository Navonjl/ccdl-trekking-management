"""Build the CCD Lab mini project report from the department template.

    uv run --with python-docx python docs/report/build_report.py

Fills the cover page, certificate and date of ccd-report-template.docx in
place, removes the template's placeholder chapters, and appends the chapters
written in report.md in the template's own formatting: Times New Roman,
chapter titles centred in bold, captions centred, tables captioned above and
figures below. LibreOffice then exports the PDF.

report.md uses a small markup:

    # CHAPTER 1 | INTRODUCTION      chapter number | chapter title
    ## 1.1 Objectives               section heading
    ### 2.3.1 Something             subsection heading
    - item                          bullet
    ![Fig. 2.1: Caption](figures/x.png){width=6.0}
    Table: Table 1.1: Caption       followed by a | pipe | table |
    EQ: expression | (2.1)          a numbered equation, centred
    [1] Reference text              a reference entry
    anything else                   a body paragraph (blank line ends it)
"""

import copy
import re
import subprocess
import sys
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

HERE = Path(__file__).resolve().parent
TEMPLATE = HERE / "ccd-report-template.docx"
SOURCE = HERE / "report.md"
OUT = HERE.parent / "ccdl-mini-project-report.docx"

TITLE = "Head in the Clouds: A Cloud-Native Trek Booking Platform on Microsoft Azure"
STUDENTS = [
    ("Navon John Lukose", "240905602"),
    ("Srishti Sudesh", "240905626"),
]
REPORT_DATE = date.today().strftime("%d/%m/%Y")

FONT = "Times New Roman"
BODY_PT = 12
CAPTION_PT = 10


# --- small helpers ------------------------------------------------------------

def set_font(run, size=None, bold=None, italic=None):
    run.font.name = FONT
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def replace_text(paragraph, new_text, size=None, bold=None):
    """Put new_text in the paragraph's first run, keeping its formatting."""
    runs = paragraph.runs
    if not runs:
        run = paragraph.add_run(new_text)
        set_font(run, size, bold)
        return
    runs[0].text = new_text
    for run in runs[1:]:
        run.text = ""


def add_inline(paragraph, text, size=BODY_PT):
    """Add text, treating `code` spans as monospace and **x** as bold."""
    for part in re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text):
        if not part:
            continue
        if part.startswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_font(run, size)
            run.font.name = "Liberation Mono"
            run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), "Liberation Mono")
            run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), "Liberation Mono")
            run.font.size = Pt(size - 1)
        elif part.startswith("**"):
            set_font(paragraph.add_run(part[2:-2]), size, bold=True)
        else:
            set_font(paragraph.add_run(part), size)


def spacing(paragraph, before=0, after=6, line=1.15):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


# --- cover, certificate --------------------------------------------------------

def fill_front_matter(doc):
    paras = doc.paragraphs
    by_text = {p.text.strip(): p for p in paras if p.text.strip()}

    replace_text(by_text["TITLE"], TITLE)
    replace_text(by_text["DD/MM/2026"], REPORT_DATE)

    # Student table: two students, so drop the two spare rows.
    table = doc.tables[0]
    for row, (name, reg) in zip(table.rows, STUDENTS):
        replace_text(row.cells[0].paragraphs[0], name)
        replace_text(row.cells[1].paragraphs[0], reg)
    for row in list(table.rows)[len(STUDENTS):]:
        row._tr.getparent().remove(row._tr)

    # Certificate: rewrite the sentence, keeping the paragraph's formatting.
    cert = next(p for p in paras if p.text.startswith("This is to certify"))
    template_run = cert.runs[0]
    for run in cert.runs:
        run._element.getparent().remove(run._element)
    names = " and ".join(f"{n} ({r})" for n, r in STUDENTS)
    pieces = [
        ("This is to certify that the project titled ", False),
        (TITLE, True),
        (" is a record of the bonafide work done by ", False),
        (names, True),
        (" submitted in partial fulfilment of the requirements for the award of the "
         "Degree of Bachelor of Technology (B.Tech.) in SCHOOL OF COMPUTER ENGINEERING "
         "of Manipal Institute of Technology, Manipal, Karnataka, (A Constituent "
         "Institute of Manipal Academy of Higher Education), during the academic year "
         "2026-2027.", False),
    ]
    for text, bold in pieces:
        run = copy.deepcopy(template_run._element)
        cert._p.append(run)
        new = cert.runs[-1]
        new.text = text
        new.bold = bold


def drop_placeholder_body(doc):
    """Delete everything from the template's 'CHAPTER 1' heading onwards."""
    body = doc.element.body
    start = next(p for p in doc.paragraphs if p.text.strip() == "CHAPTER 1")
    started = False
    for el in list(body):
        if el is start._p:
            started = True
        if started and el.tag != qn("w:sectPr"):
            body.remove(el)


# --- body ---------------------------------------------------------------------

def page_break(doc):
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)
    spacing(p, 0, 0)


def chapter(doc, number, title, first):
    if not first:
        page_break(doc)
    for text in (number, title):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_font(p.add_run(text), 14, bold=True)
        spacing(p, 0, 2)
    doc.paragraphs[-1].paragraph_format.space_after = Pt(14)


def heading(doc, text, level):
    p = doc.add_paragraph()
    set_font(p.add_run(text), 13 if level == 2 else BODY_PT, bold=True)
    spacing(p, 10 if level == 2 else 8, 4)
    p.paragraph_format.keep_with_next = True


def body_paragraph(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    add_inline(p, text)
    spacing(p, 0, 6, 1.3)


def bullet(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    fmt = p.paragraph_format
    fmt.left_indent = Inches(0.35)
    fmt.first_line_indent = Inches(-0.2)
    add_inline(p, "•  " + text)
    spacing(p, 0, 3, 1.25)


def caption(doc, text, keep_with_next=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run(text), CAPTION_PT, bold=True)
    spacing(p, 4, 10 if not keep_with_next else 4)
    p.paragraph_format.keep_with_next = keep_with_next


def figure(doc, cap, path, width):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(HERE / path), width=Inches(width))
    spacing(p, 6, 0, 1.0)
    caption(doc, cap)


def table(doc, cap, rows):
    caption(doc, cap, keep_with_next=True)
    header, *data = rows
    t = doc.add_table(rows=len(rows), cols=len(header))
    t.style = "Table Grid" if "Table Grid" in [s.name for s in doc.styles] else None
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r, cells in enumerate(rows):
        for c, text in enumerate(cells):
            cell = t.cell(r, c)
            para = cell.paragraphs[0]
            add_inline(para, text, size=10)
            spacing(para, 1, 1, 1.0)
            if r == 0:
                for run in para.runs:
                    run.bold = True
    borders(t)
    after = doc.add_paragraph()
    spacing(after, 0, 6)


def borders(t):
    """Single hairline borders on every cell, like the template's tables."""
    tbl_pr = t._tbl.tblPr
    el = tbl_pr.find(qn("w:tblBorders"))
    if el is None:
        el = tbl_pr.makeelement(qn("w:tblBorders"), {})
        tbl_pr.append(el)
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b = el.makeelement(qn(f"w:{side}"), {qn("w:val"): "single", qn("w:sz"): "4",
                                            qn("w:space"): "0", qn("w:color"): "000000"})
        el.append(b)


def equation(doc, expr, number):
    """Expression centred on the line, its number against the right margin."""
    from docx.enum.text import WD_TAB_ALIGNMENT
    p = doc.add_paragraph()
    section = doc.sections[0]
    width = section.page_width - section.left_margin - section.right_margin
    stops = p.paragraph_format.tab_stops
    stops.add_tab_stop(int(width / 2), WD_TAB_ALIGNMENT.CENTER)
    stops.add_tab_stop(width, WD_TAB_ALIGNMENT.RIGHT)
    set_font(p.add_run("\t" + expr), BODY_PT, italic=True)
    set_font(p.add_run("\t" + number), BODY_PT)
    spacing(p, 4, 8, 1.2)


def reference(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT  # justifying long URLs leaves wide gaps
    fmt = p.paragraph_format
    fmt.left_indent = Inches(0.45)
    fmt.first_line_indent = Inches(-0.45)
    add_inline(p, text, size=11)
    spacing(p, 0, 4, 1.15)


def render_body(doc, source):
    lines = source.splitlines()
    i, first_chapter, para = 0, True, []

    def flush():
        if para:
            body_paragraph(doc, " ".join(para))
            para.clear()

    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip():
            flush()
        elif line.startswith("# "):
            flush()
            number, title = [s.strip() for s in line[2:].split("|")]
            chapter(doc, number, title, first_chapter)
            first_chapter = False
        elif line.startswith("## "):
            flush()
            heading(doc, line[3:], 2)
        elif line.startswith("### "):
            flush()
            heading(doc, line[4:], 3)
        elif line.startswith("- "):
            flush()
            bullet(doc, line[2:])
        elif line.startswith("!["):
            flush()
            m = re.match(r"!\[(.+?)\]\((.+?)\)(?:\{width=([\d.]+)\})?", line)
            figure(doc, m.group(1), m.group(2), float(m.group(3) or 6.0))
        elif line.startswith("Table: "):
            flush()
            cap = line[len("Table: "):]
            rows = []
            i += 1
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-+:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            table(doc, cap, rows)
            continue
        elif line.startswith("EQ: "):
            flush()
            expr, number = [s.strip() for s in line[4:].split("|")]
            equation(doc, expr, number)
        elif re.match(r"\[\d+\] ", line):
            flush()
            reference(doc, line)
        else:
            para.append(line.strip())
        i += 1
    flush()


def main():
    doc = Document(TEMPLATE)
    fill_front_matter(doc)
    drop_placeholder_body(doc)
    render_body(doc, SOURCE.read_text())
    doc.save(OUT)
    print(f"wrote {OUT}")
    if "--no-pdf" not in sys.argv:
        subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir",
                        str(OUT.parent), str(OUT)], check=True, capture_output=True)
        print(f"wrote {OUT.with_suffix('.pdf')}")


if __name__ == "__main__":
    main()
